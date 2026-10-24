"""Relatório mensal em Excel e PDF (receita, a receber, projetos e solicitações)."""
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .config import OUT_DIR

BRL = '"R$" #,##0.00'
STATUS_PROJETO = {"proposta": "Proposta", "em_andamento": "Em andamento", "pausado": "Pausado", "concluido": "Concluído", "cancelado": "Cancelado"}


def brl(value) -> str:
    return f"R$ {float(value or 0):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def month_bounds(month: str) -> tuple[date, date]:
    """'2026-10' -> (2026-10-01, 2026-11-01)"""
    year, mon = (int(p) for p in month.split("-"))
    start = date(year, mon, 1)
    return start, date(year + (mon == 12), mon % 12 + 1, 1)


def collect(conn, month: str) -> dict:
    start, end = month_bounds(month)
    q = lambda sql, params=(): conn.execute(sql, params).fetchall()
    received = q(
        """SELECT y.paid_at, y.description, y.amount, p.title AS project, c.name AS client
           FROM payments y JOIN projects p ON p.id = y.project_id LEFT JOIN clients c ON c.id = p.client_id
           WHERE y.paid_at >= %s AND y.paid_at < %s ORDER BY y.paid_at""", (start, end))
    open_payments = q(
        """SELECT y.due_date, y.description, y.amount, p.title AS project, c.name AS client,
                  (y.due_date < CURRENT_DATE) AS overdue
           FROM payments y JOIN projects p ON p.id = y.project_id LEFT JOIN clients c ON c.id = p.client_id
           WHERE y.paid_at IS NULL ORDER BY y.due_date""")
    projects = q(
        """SELECT p.title, p.status, p.stage, p.progress, p.budget, p.due_date, c.name AS client
           FROM projects p LEFT JOIN clients c ON c.id = p.client_id
           WHERE p.status IN ('proposta', 'em_andamento', 'pausado') OR (p.created_at >= %s AND p.created_at < %s)
           ORDER BY p.due_date NULLS LAST""", (start, end))
    counts = q(
        """SELECT (SELECT count(*) FROM clients WHERE created_at >= %s AND created_at < %s) AS new_clients,
                  (SELECT count(*) FROM leads WHERE created_at >= %s AND created_at < %s) AS new_leads,
                  (SELECT count(*) FROM projects WHERE status = 'concluido' AND created_at < %s) AS done_projects""",
        (start, end, start, end, end))[0]
    return {
        "month": month, "received": received, "open": open_payments, "projects": projects,
        "total_received": sum(r["amount"] for r in received),
        "total_open": sum(r["amount"] for r in open_payments),
        "total_overdue": sum(r["amount"] for r in open_payments if r["overdue"]),
        **{k: int(v) for k, v in counts.items()},
    }


def _sheet(ws, headers, rows, widths, money_cols=(), date_cols=()):
    ws.append(headers)
    fill = PatternFill("solid", fgColor="0F2947")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        ws.append(row)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    for col in money_cols:
        for cell in ws[chr(64 + col)][1:]:
            cell.number_format = BRL
    for col in date_cols:
        for cell in ws[chr(64 + col)][1:]:
            cell.number_format = "DD/MM/YYYY"
    ws.freeze_panes = "A2"


def build_xlsx(data: dict, path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    ws.append([f"Relatório Carcará · {data['month']}"])
    ws["A1"].font = Font(bold=True, size=14)
    for label, value, fmt in [
        ("Recebido no mês", float(data["total_received"]), BRL),
        ("A receber (total em aberto)", float(data["total_open"]), BRL),
        ("Em atraso", float(data["total_overdue"]), BRL),
        ("Novos clientes", data["new_clients"], "0"),
        ("Novas solicitações", data["new_leads"], "0"),
        ("Projetos concluídos (acumulado)", data["done_projects"], "0"),
    ]:
        ws.append([label, value])
        ws.cell(ws.max_row, 2).number_format = fmt
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 18

    _sheet(wb.create_sheet("Recebimentos"), ["Data", "Descrição", "Projeto", "Cliente", "Valor"],
           [[r["paid_at"], r["description"], r["project"], r["client"], float(r["amount"])] for r in data["received"]],
           [13, 34, 30, 26, 16], money_cols=(5,), date_cols=(1,))
    _sheet(wb.create_sheet("A receber"), ["Vencimento", "Descrição", "Projeto", "Cliente", "Valor", "Situação"],
           [[r["due_date"], r["description"], r["project"], r["client"], float(r["amount"]), "Atrasado" if r["overdue"] else "Pendente"] for r in data["open"]],
           [13, 34, 30, 26, 16, 12], money_cols=(5,), date_cols=(1,))
    _sheet(wb.create_sheet("Projetos"), ["Projeto", "Cliente", "Status", "Progresso (%)", "Orçamento", "Prazo"],
           [[r["title"], r["client"], STATUS_PROJETO.get(r["status"], r["status"]), r["progress"], float(r["budget"]), r["due_date"]] for r in data["projects"]],
           [34, 26, 16, 14, 16, 13], money_cols=(5,), date_cols=(6,))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def build_pdf(data: dict, path: Path) -> Path:
    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm, bottomMargin=2 * cm,
                            title=f"Relatório Carcará {data['month']}", author="Carcará Tecnologia")
    head = colors.HexColor("#0F2947")

    def table(rows, widths):
        t = Table(rows, colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), head), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTSIZE", (0, 0), (-1, -1), 8.5), ("GRID", (0, 0), (-1, -1), .25, colors.lightgrey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F6F9")]),
            ("ALIGN", (-1, 0), (-1, -1), "RIGHT"),
        ]))
        return t

    story = [Paragraph(f"Relatório mensal · {data['month']}", styles["Title"]), Paragraph("Carcará Tecnologia", styles["Normal"]), Spacer(1, 14)]
    story.append(table([
        ["Indicador", "Valor"],
        ["Recebido no mês", brl(data["total_received"])],
        ["A receber (em aberto)", brl(data["total_open"])],
        ["Em atraso", brl(data["total_overdue"])],
        ["Novos clientes", str(data["new_clients"])],
        ["Novas solicitações", str(data["new_leads"])],
    ], [10 * cm, 7 * cm]))

    story += [Spacer(1, 16), Paragraph("Recebimentos do mês", styles["Heading2"])]
    story.append(table([["Data", "Descrição", "Cliente", "Valor"]] + [
        [r["paid_at"].strftime("%d/%m/%Y"), Paragraph(r["description"], styles["BodyText"]), r["client"] or "—", brl(r["amount"])] for r in data["received"]
    ], [2.3 * cm, 7 * cm, 4.7 * cm, 3 * cm]) if data["received"] else Paragraph("Nenhum recebimento no período.", styles["Normal"]))

    story += [Spacer(1, 16), Paragraph("Valores em aberto", styles["Heading2"])]
    story.append(table([["Vencimento", "Descrição", "Situação", "Valor"]] + [
        [r["due_date"].strftime("%d/%m/%Y"), Paragraph(r["description"], styles["BodyText"]), "Atrasado" if r["overdue"] else "Pendente", brl(r["amount"])] for r in data["open"]
    ], [2.5 * cm, 7.5 * cm, 3 * cm, 4 * cm]) if data["open"] else Paragraph("Nada em aberto.", styles["Normal"]))

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.build(story)
    return path


def generate(conn, month: str, out_dir: Path = OUT_DIR) -> tuple[dict, list[Path]]:
    data = collect(conn, month)
    files = [build_xlsx(data, out_dir / f"relatorio-{month}.xlsx"), build_pdf(data, out_dir / f"relatorio-{month}.pdf")]
    return data, files
