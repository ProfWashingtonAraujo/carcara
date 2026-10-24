from datetime import date, datetime

from openpyxl import load_workbook

from carcara import alerts, reports


def empty():
    return {"overdue": [], "due_soon": [], "projects": [], "leads": []}


def test_no_message_when_nothing_pending():
    assert alerts.build_message(empty()) is None


def test_message_lists_sections():
    data = empty()
    data["overdue"] = [{"description": "Entrada", "amount": 1500, "due_date": date(2026, 10, 1), "project": "Site", "client": "Ana"}]
    data["leads"] = [{"name": "Beto", "service": None, "created_at": datetime(2026, 10, 20)}]
    subject, body = alerts.build_message(data)
    assert "2 pendência" in subject
    assert "R$ 1.500,00" in body and "EM ATRASO" in body and "Beto" in body


def test_month_bounds_handles_december():
    assert reports.month_bounds("2026-12") == (date(2026, 12, 1), date(2027, 1, 1))
    assert reports.month_bounds("2026-02") == (date(2026, 2, 1), date(2026, 3, 1))


def test_brl_format():
    assert reports.brl(1234.5) == "R$ 1.234,50"
    assert reports.brl(None) == "R$ 0,00"


def sample():
    return {
        "month": "2026-10", "total_received": 3000, "total_open": 500, "total_overdue": 200,
        "new_clients": 2, "new_leads": 5, "done_projects": 1,
        "received": [{"paid_at": date(2026, 10, 5), "description": "Entrada <50%>", "amount": 3000, "project": "Site", "client": "Ana"}],
        "open": [{"due_date": date(2026, 11, 5), "description": "Parcela", "amount": 500, "project": "Site", "client": None, "overdue": False}],
        "projects": [{"title": "Site", "client": "Ana", "status": "em_andamento", "progress": 40, "budget": 6000, "due_date": None}],
    }


def test_xlsx_and_pdf_are_generated(tmp_path):
    x = reports.build_xlsx(sample(), tmp_path / "r.xlsx")
    wb = load_workbook(x)
    assert wb.sheetnames == ["Resumo", "Recebimentos", "A receber", "Projetos"]
    assert wb["Resumo"]["B2"].value == 3000
    p = reports.build_pdf(sample(), tmp_path / "r.pdf")
    assert p.read_bytes().startswith(b"%PDF")
